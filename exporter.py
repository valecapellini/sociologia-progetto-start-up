import csv
import logging
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

logger = logging.getLogger(__name__)


def export_to_excel(startups: list[dict], region: str = "liguria", output_dir: str = ".") -> str:
    """Esporta la lista di startup in un file Excel formattato."""
    if not startups:
        logger.warning("Nessuna startup da esportare.")
        return ""

    Path(output_dir).mkdir(parents=True, exist_ok=True)

    region_title = region.strip().title()

    wb = Workbook()
    ws = wb.active
    ws.title = f"Startup {region_title}"

    # Raccogliere tutti i campi presenti (preservando ordine)
    all_fields = []
    seen = set()
    for s in startups:
        for key in s:
            if key not in seen:
                all_fields.append(key)
                seen.add(key)

    # Header
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    for col_idx, field in enumerate(all_fields, 1):
        cell = ws.cell(row=1, column=col_idx, value=field)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Dati
    for row_idx, startup in enumerate(startups, 2):
        for col_idx, field in enumerate(all_fields, 1):
            value = startup.get(field, "")
            ws.cell(row=row_idx, column=col_idx, value=value)

    # Auto-larghezza colonne
    for col_idx, field in enumerate(all_fields, 1):
        max_len = len(str(field))
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(max_len + 2, 60)

    # Salva
    today = datetime.now().strftime("%Y%m%d")
    region_slug = region.strip().lower().replace(" ", "_").replace("'", "")
    filename = f"{output_dir}/startup_{region_slug}_{today}.xlsx"
    wb.save(filename)
    logger.info(f"File Excel salvato: {filename} ({len(startups)} startup)")

    # CSV
    csv_filename = filename.replace(".xlsx", ".csv")
    with open(csv_filename, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=all_fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(startups)
    logger.info(f"File CSV salvato: {csv_filename}")

    return filename
